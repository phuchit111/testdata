"""Read-only structural audit for the FruitBlend24 workflow design.

Run with the bundled Python runtime. Does not modify or export the source workbook.
The SKU mapping below is a candidate mapping for planning, not a released clean dataset.
"""
from pathlib import Path
import hashlib
import json
import sys

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data/raw/FruitBlend24_Intern_Case_Data.xlsx"
sys.stdout.reconfigure(encoding="utf-8")
workbook = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)


def table(name):
    rows = workbook[name].values
    headers = next(rows)
    return pd.DataFrame(rows, columns=headers)


orders = table("orders_hourly")
aliases = {
    "Watermelon": ["Watermelon Ice", "Watermelon smoothy", "Watermelon Smoothie", "watermelon ice", "แตงโมปั่นสด", "แตงโมปั่น", "แตงโมปั่น (ไซส์ M)"],
    "Pineapple": ["สับปะรดปั่น", "Pineapple Ice", "pineapple ice", "สัปปะรดปั่น", "Pineapple Smoothie", "Pineapple smoothie "],
    "Guava": ["Guava Ice", "Guava Smoothie", "ฝรั่งปั่น (ใส่เกลือ)", "guava ice", "ฝรั่งปั่น"],
    "PassionFruit": ["Passion Fruit Smoothie", "เสาวรสปั่น", "Passion Fruit Ice", "เสาวรสปั่นเข้มข้น", "passionfruit ice"],
    "MixedBerryPremium": ["เบอร์รี่ปั่นพรีเมียม", "berry premium smoothie", "Mixed Berry (Premium)", "เบอร์รี่รวมพรีเมียม", "Mixed Berry Premium"],
}
mapping = {alias: sku for sku, names in aliases.items() for alias in names}
dedup = orders.drop_duplicates().copy()
dedup["sku"] = dedup.item_name_raw.map(mapping)
unmapped = dedup.loc[dedup.sku.isna()]
candidate = dedup.loc[dedup.sku.notna()].copy()
candidate["date"] = pd.to_datetime(candidate.date)
candidate["week_start"] = candidate.date - pd.to_timedelta(candidate.date.dt.dayofweek, unit="D")
code = candidate.rate_code.str.extract(r"^(LM|GB)-(RC\d{3})(?:-([A-Z]{2}))?$")
code.columns = ["platform_code", "base_code", "sku_code"]
candidate = pd.concat([candidate, code], axis=1)
fruit = table("weekly_fruit_cost")
fruit["week_start"] = pd.to_datetime(fruit.week_start)
joined = candidate.merge(fruit, on=["week_start", "sku"], how="left", validate="many_to_one")
missing_cost = joined.loc[joined.fruit_cost_per_cup_thb.isna()]
rates = table("rate_code_dim").set_index("base_code").discount_pct
sku_codes = table("sku_code_dim").set_index("sku_code").sku
assumptions = table("cost_assumptions").iloc[:5].set_index("sku")
expected_price = candidate.sku.map(assumptions.base_price_thb) * (1 - candidate.base_code.map(rates))
sku_suffix_mismatch = candidate.sku_code.notna() & (candidate.sku != candidate.sku_code.map(sku_codes))
waste = table("waste_daily")
waste["date"] = pd.to_datetime(waste.date)
waste["week_start"] = waste.date - pd.to_timedelta(waste.date.dt.dayofweek, unit="D")
waste_cost = waste.merge(fruit, on=["week_start", "sku"], how="left", validate="many_to_one")
expected_waste_cost = waste_cost.units_wasted * (waste_cost.fruit_cost_per_cup_thb + waste_cost.sku.map(assumptions.packaging_cost_per_cup_thb))
dates = pd.to_datetime(orders.date, errors="coerce")
report = {
    "scope": "Read-only planning audit; candidate normalization is not a final cleaned data release or business analysis.",
    "source": str(SOURCE),
    "source_sha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
    "profile_script_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
    "sheets": {s.title: {"data_rows_including_internal_blanks": s.max_row - 1, "columns": s.max_column} for s in workbook},
    "raw_orders": {
        "rows": len(orders),
        "period_start": str(dates.min().date()),
        "period_end": str(dates.max().date()),
        "distinct_raw_names": int(orders.item_name_raw.nunique()),
        "exact_duplicate_excess_rows": int(orders.duplicated().sum()),
        "null_cells_by_column": orders.isna().sum().to_dict(),
        "invalid_dates": int(dates.isna().sum()),
        "invalid_hours": int((~orders.hour.between(0, 23)).sum()),
        "nonpositive_units": int((orders.units_sold <= 0).sum()),
        "revenue_identity_mismatch_rows_over_001_thb": int(((orders.units_sold * orders.price_thb - orders.gross_revenue_thb).abs() > 0.01).sum()),
        "weekday_mismatch_rows": int((dates.dt.dayofweek != orders.day_of_week).sum()),
    },
    "candidate_preparation": {
        "rows_after_exact_dedup": len(dedup),
        "unmapped_rows_after_dedup": len(unmapped),
        "unmapped_names_after_dedup": unmapped.item_name_raw.value_counts().to_dict(),
        "retained_menu_rows": len(candidate),
        "duplicate_excess_at_date_hour_kitchen_sku": int(candidate.duplicated(["date", "hour", "kitchen", "sku"]).sum()),
        "invalid_rate_format_rows": int(candidate.platform_code.isna().sum()),
        "unknown_base_codes": sorted(set(candidate.base_code) - set(rates.index)),
        "sku_suffix_mismatch_rows": int(sku_suffix_mismatch.sum()),
        "price_discount_formula_mismatch_rows_over_001_thb": int(((candidate.price_thb - expected_price).abs() > 0.01).sum()),
        "missing_exact_week_fruit_cost_rows": len(missing_cost),
        "missing_exact_week_fruit_cost_keys": missing_cost[["date", "week_start", "sku"]].drop_duplicates().astype(str).to_dict("records"),
        "sku_first_date": candidate.groupby("sku").date.min().dt.strftime("%Y-%m-%d").to_dict(),
        "sku_last_date": candidate.groupby("sku").date.max().dt.strftime("%Y-%m-%d").to_dict(),
        "candidate_alias_mapping": aliases,
    },
    "waste": {
        "rows": len(waste),
        "missing_estimated_waste_cost_rows": int(waste.est_waste_cost_thb.isna().sum()),
        "units_with_missing_estimated_waste_cost": int(waste.loc[waste.est_waste_cost_thb.isna(), "units_wasted"].sum()),
        "duplicate_excess_at_date_kitchen_sku": int(waste.duplicated(["date", "kitchen", "sku"]).sum()),
        "missing_exact_week_fruit_cost_rows": int(waste_cost.fruit_cost_per_cup_thb.isna().sum()),
        "recomputed_waste_cost_mismatch_rows_over_001_thb_where_cost_available": int(((waste_cost.est_waste_cost_thb - expected_waste_cost).abs() > 0.01).sum()),
    },
    "limitations": [
        "Positive sales observations do not prove that absent hourly rows are zero demand or that no stockouts occurred.",
        "No customer/order IDs, ingredient recipes, stock balances, supplier lead times, shelf life, or purchase orders are present in the provided tables.",
        "Monthly budget does not define the gross-profit cost boundary; confirm or disclose the adopted definition.",
        "Observed selling prices already include the applied discount; do not deduct it again.",
        "Waste estimate includes fruit and packaging only, as specified in data_dictionary.",
    ],
}
workbook.close()
destination = ROOT / "planning/data_readiness.json"
destination.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
if "--quiet" not in sys.argv:
    print(json.dumps({
        "report_path": str(destination),
        "raw_rows": len(orders),
        "duplicates_removed_in_candidate": int(orders.duplicated().sum()),
        "non_menu_rows_in_candidate": len(unmapped),
        "retained_candidate_rows": len(candidate),
        "sales_rows_missing_exact_week_cost": len(missing_cost),
        "waste_rows_missing_estimated_cost": int(waste.est_waste_cost_thb.isna().sum()),
    }, ensure_ascii=False))
